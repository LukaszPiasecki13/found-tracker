"""XTB account-report adapter (`.xlsx`) for the `ImportParser` port.

Operations come from the `Cash Operations` sheet (the source of truth), plus
the stock splits: XTB books a split as a position transfer that has no cash
row, so it is derived from `Closed Positions` and `Open Positions`, which are
also read for the expectations the imported operations are verified against.
`openpyxl` is imported only here.

The archive is checked *before* it is opened (zip/XML bombs): total uncompressed
size, entry count and compression ratio are capped, and so is the row count.
"""

import re
import zipfile
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.core.import_parser import (
    ImportExpectations,
    ImportParseError,
    ParsedRow,
    ParseIssue,
    ParseResult,
)

MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 100
MAX_COMPRESSION_RATIO = 100
MAX_ROWS = 20_000

_CASH_SHEET = "Cash Operations"
_OPEN_SHEET = "Open Positions"
_CLOSED_SHEET = "Closed Positions"

_TOTAL = "total"
_TYPE_MAP = {
    "deposit": "deposit",
    "withdrawal": "withdrawal",
    "stock purchase": "buy",
    "stock sell": "sell",
    "dividend": "dividend",
}
# Cash that is neither a deposit nor a trade: interest, taxes, fees, the result
# of a CFD. Booked as income (`interest`) or cost (`fee`) by the sign of the
# amount, without an asset; the source's type and ticker go into the notes.
_INCOME_COST_TYPES = frozenset(
    {
        "free funds interest",
        "free funds interest tax",
        "withholding tax",
        "sec fee",
        "close trade",
        "swap",
        "correction",
    }
)
_CLASS_NAMES = {"STOCK": "Stock", "ETF": "ETF"}
_TRADE_TYPES = frozenset({"buy", "sell"})
_ASSET_TYPES = frozenset({"buy", "sell", "dividend"})
# Whether the source books the amount as positive (cash in) for a type.
_EXPECTED_SIGN = {
    "deposit": True,
    "withdrawal": False,
    "buy": False,
    "sell": True,
    "dividend": True,
}
# "OPEN BUY 84 @ 40.710", "CLOSE BUY 40/84 @ 40.730", "CLOSE BUY 0.67 @ 495.00"
_TRADE_COMMENT = re.compile(
    r"^\s*(?P<phase>OPEN|CLOSE)\s+(?:BUY|SELL)\s+"
    r"(?P<quantity>\d+(?:\.\d+)?)(?:\s*/\s*\d+(?:\.\d+)?)?\s*@\s*"
    r"(?P<price>\d+(?:\.\d+)?)",
    re.IGNORECASE,
)
_EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=UTC)
# A split is booked as a transfer: the old positions are closed with a comment
# "... Transfer Out" and the new ones open within hours, `ratio` times the volume.
_TRANSFER_OUT = "transfer out"
_TRANSFER_WINDOW = timedelta(hours=6)
_RATIO_PLACES = Decimal("0.000001")

type _Cells = dict[str, Any]


def _check_archive(content: bytes) -> None:
    """Raises `ImportParseError` unless `content` is a zip within the limits."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
    except zipfile.BadZipFile as err:
        raise ImportParseError("The file is not a valid .xlsx archive") from err
    uncompressed = sum(info.file_size for info in infos)
    compressed = sum(info.compress_size for info in infos) or 1
    if (
        len(infos) > MAX_ARCHIVE_ENTRIES
        or uncompressed > MAX_UNCOMPRESSED_BYTES
        or uncompressed / compressed > MAX_COMPRESSION_RATIO
    ):
        raise ImportParseError("The file exceeds the allowed archive limits")


def _decimal(value: object) -> Decimal | None:
    """Cell number -> `Decimal` via `str` (never `Decimal(float)`, ADR-0010)."""
    if value is None or isinstance(value, bool) or value == "":
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def _money(value: object) -> Decimal | None:
    """A sheet total: the report's float noise (4.03999999999985) rounded to
    cents."""
    number = _decimal(value)
    return None if number is None else number.quantize(Decimal("0.01"))


def _timestamp(value: object) -> datetime | None:
    """A cell holding a date (as `openpyxl` returns it) or an Excel serial
    (1900 system) -> aware UTC datetime."""
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    number = None if isinstance(value, bool) else _decimal(value)
    if number is None:
        return None
    return (_EXCEL_EPOCH + timedelta(days=float(number))).replace(microsecond=0)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_rows(workbook: Any, name: str) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """(1-based sheet row number, cells) of a sheet, if it exists."""
    if name not in workbook.sheetnames:
        return
    sheet = workbook[name]
    # XTB writes a wrong sheet dimension (`A1`); a read-only sheet would trust it
    # and stop after the first row.
    sheet.reset_dimensions()
    for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
        if number > MAX_ROWS:
            raise ImportParseError(f"Sheet '{name}' has more than {MAX_ROWS} rows")
        yield number, cells


def _table(
    workbook: Any,
    name: str,
    first_header: str,
    also: str | None = None,
    required: bool = False,
) -> Iterator[tuple[int, _Cells]]:
    """Rows of a sheet below its header row (the one starting with
    `first_header` and, if given, also holding the column `also`), as
    {header: cell}; rows above the header are skipped."""
    header: list[str] | None = None
    for number, cells in _sheet_rows(workbook, name):
        if header is None:
            names = [_text(cell) for cell in cells]
            if names[:1] == [first_header] and (also is None or also in names):
                header = names
            continue
        yield number, dict(zip(header, cells, strict=False))
    if header is None and required:
        raise ImportParseError(f"Header row not found in sheet '{name}'")


@dataclass(frozen=True, slots=True)
class _Lot:
    """A position (or the part of it a closing row covers), as the Closed and
    Open Positions sheets list it."""

    ticker: str
    category: str
    volume: Decimal
    opened: datetime
    closed: datetime | None
    position_id: str
    transfer_out: bool


def _closed_lot(cells: _Cells) -> _Lot | None:
    ticker = _text(cells.get("Ticker"))
    volume = _decimal(cells.get("Volume"))
    opened = _timestamp(cells.get("Open Time (UTC)"))
    closed = _timestamp(cells.get("Close Time (UTC)"))
    if not ticker or volume is None or volume <= 0 or opened is None or closed is None:
        return None
    return _Lot(
        ticker=ticker,
        category=_text(cells.get("Category")),
        volume=volume,
        opened=opened,
        closed=closed,
        position_id=_text(cells.get("Position ID")),
        transfer_out=_TRANSFER_OUT in _text(cells.get("Comment")).lower(),
    )


def _open_lot(cells: _Cells) -> _Lot | None:
    ticker = _text(cells.get("Ticker"))
    volume = _decimal(cells.get("Volume"))
    opened = _timestamp(cells.get("Open time (UTC)"))
    if (
        _text(cells.get("Type")).upper() != "BUY"
        or not ticker
        or volume is None
        or volume <= 0
        or opened is None
    ):
        return None
    return _Lot(
        ticker=ticker,
        category=_text(cells.get("Category")),
        volume=volume,
        opened=opened,
        closed=None,
        position_id=_text(cells.get("Instrument/Position")),
        transfer_out=False,
    )


class _RowError(Exception):
    """A row that cannot become an operation; becomes a `ParseIssue`."""


def _notes(source_type: str, cells: _Cells) -> str:
    """What the user sees of an income/cost row: the source's type, its ticker
    and its comment."""
    parts = (source_type, _text(cells.get("Ticker")), _text(cells.get("Comment")))
    return " | ".join(part for part in parts if part)


def _is_blank(row: _Cells) -> bool:
    return all(_text(cell) == "" for cell in row.values())


class XtbParser:
    """Reads an XTB "account report" workbook."""

    parser_id = "xtb"

    def sniff(self, filename: str, content: bytes) -> bool:
        if not filename.lower().endswith(".xlsx"):
            return False
        try:
            _check_archive(content)
            with zipfile.ZipFile(BytesIO(content)) as archive:
                workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        except ImportParseError, KeyError, UnicodeDecodeError:
            return False
        return f'name="{_CASH_SHEET}"' in workbook_xml

    def parse(self, content: bytes) -> ParseResult:
        _check_archive(content)
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as err:  # openpyxl raises many types for corrupt input
            raise ImportParseError("The file cannot be read as an XTB report") from err
        try:
            if _CASH_SHEET not in workbook.sheetnames:
                raise ImportParseError(f"Sheet '{_CASH_SHEET}' not found")
            rows, issues, cash_total = self._cash_operations(workbook)
            split_rows, split_issues = self._splits(workbook, rows, issues)
            rows += split_rows
            issues += split_issues
            expectations = ImportExpectations(
                open_positions=self._open_positions(workbook),
                cash_total=cash_total,
                closed_profit=self._closed_profit(workbook),
            )
        finally:
            workbook.close()
        return ParseResult(
            rows=_stable_refs(rows),
            issues=issues,
            expectations=expectations,
        )

    # --- Cash Operations ---

    def _cash_operations(
        self, workbook: Any
    ) -> tuple[list[ParsedRow], list[ParseIssue], Decimal | None]:
        """Rows, issues and the sheet's Total."""
        rows: list[ParsedRow] = []
        issues: list[ParseIssue] = []
        cash_total: Decimal | None = None
        for number, cells in _table(workbook, _CASH_SHEET, "Type", required=True):
            if _is_blank(cells):
                continue
            source_type = _text(cells.get("Type"))
            if source_type.lower() == _TOTAL:
                cash_total = _money(cells.get("Amount"))
            else:
                try:
                    rows.append(self._row(number, cells))
                except _RowError as err:
                    raw = {key: _text(value) for key, value in cells.items()}
                    issues.append(ParseIssue(number, str(err), raw))
        return rows, issues, cash_total

    def _row(self, number: int, cells: _Cells) -> ParsedRow:
        source_type = _text(cells.get("Type"))
        income_cost = source_type.lower() in _INCOME_COST_TYPES
        mapped = _TYPE_MAP.get(source_type.lower())
        if mapped is None and not income_cost:
            raise _RowError(f"Unknown operation type '{source_type}'")
        amount = _money(cells.get("Amount"))
        when = _timestamp(cells.get("Time"))
        if amount is None or when is None:
            raise _RowError("Missing or unreadable amount or time")
        operation_type = mapped or ""
        if income_cost:
            operation_type = "interest" if amount > 0 else "fee"
        elif amount != 0 and (amount > 0) != _EXPECTED_SIGN[operation_type]:
            raise _RowError(f"Unexpected sign of the amount for '{source_type}'")
        cash_id = _text(cells.get("ID"))
        if not cash_id:
            raise _RowError("Missing operation ID")
        row = ParsedRow(
            row_number=number,
            operation_type=operation_type,
            operation_date=when,
            amount=abs(amount),
            external_ref=f"xtb:cash:{cash_id}",
            notes=_notes(source_type, cells)
            if income_cost
            else _text(cells.get("Comment")),
        )
        if operation_type in _TRADE_TYPES:
            row = self._trade(row, cells)
        return self._asset(row, cells) if operation_type in _ASSET_TYPES else row

    def _asset(self, row: ParsedRow, cells: _Cells) -> ParsedRow:
        """The asset a row concerns: the source's ticker, its exchange suffix
        and the asset class named by the Category column."""
        ticker = _text(cells.get("Ticker"))
        if not ticker:
            raise _RowError("An operation on an asset needs a ticker")
        category = _text(cells.get("Category"))
        return replace(
            row,
            ticker=ticker,
            exchange_hint=ticker.rpartition(".")[2].upper() if "." in ticker else None,
            asset_class=_CLASS_NAMES.get(category.upper(), category.title() or None),
        )

    def _trade(self, row: ParsedRow, cells: _Cells) -> ParsedRow:
        """Quantity and price from the comment; the id is the position's."""
        match = _TRADE_COMMENT.match(row.notes)
        position_id = _text(cells.get("Position ID"))
        if match is None or not position_id:
            raise _RowError(
                "A trade needs a Position ID and a comment like 'OPEN BUY 84 @ 40.710'"
            )
        return replace(
            row,
            quantity=Decimal(match["quantity"]),
            price=Decimal(match["price"]),
            external_ref=f"xtb:pos:{position_id}:{match['phase'].lower()}",
        )

    # --- Splits ---

    def _splits(
        self, workbook: Any, rows: list[ParsedRow], issues: list[ParseIssue]
    ) -> tuple[list[ParsedRow], list[ParseIssue]]:
        """The splits found in the position sheets: the ratio is the volume of
        the positions opened by the transfer over the volume it closed. A
        transfer that keeps the volume is no split. Their rows are numbered
        after the last Cash Operations row."""
        opened_by_cash = {
            row.external_ref.split(":")[2]
            for row in rows
            if row.external_ref.startswith("xtb:pos:")
            and row.external_ref.endswith(":open")
        }
        lots = [
            lot
            for _, cells in _table(workbook, _CLOSED_SHEET, "Instrument")
            if (lot := _closed_lot(cells)) is not None
        ]
        lots += [
            lot
            for _, cells in _table(workbook, _OPEN_SHEET, "Product", also="Volume")
            if (lot := _open_lot(cells)) is not None
        ]
        transfers: dict[tuple[str, date], list[_Lot]] = {}
        for lot in lots:
            if lot.transfer_out and lot.closed is not None:
                transfers.setdefault((lot.ticker, lot.closed.date()), []).append(lot)

        numbered = [row.row_number for row in rows] + [i.row_number for i in issues]
        last_row = max(numbered, default=0)
        split_rows: list[ParsedRow] = []
        split_issues: list[ParseIssue] = []
        for (ticker, _), old in sorted(transfers.items(), key=lambda t: t[0][1]):
            moved = min(lot.closed for lot in old if lot.closed is not None)
            old_volume = sum((lot.volume for lot in old), Decimal(0))
            new_volume = sum(
                (
                    lot.volume
                    for lot in lots
                    if lot.ticker == ticker
                    and not lot.transfer_out
                    and lot.position_id not in opened_by_cash
                    and moved < lot.opened <= moved + _TRANSFER_WINDOW
                ),
                Decimal(0),
            )
            number = last_row + len(split_rows) + len(split_issues) + 1
            if new_volume == 0:
                message = (
                    f"{ticker}: a position transfer (split) was found but not the "
                    "positions it opened - the ratio is unknown"
                )
                split_issues.append(ParseIssue(number, message, {"Ticker": ticker}))
                continue
            ratio = (new_volume / old_volume).quantize(_RATIO_PLACES)
            if ratio == 1:
                continue
            row = ParsedRow(
                row_number=number,
                operation_type="split",
                operation_date=moved,
                amount=Decimal(0),
                external_ref=f"xtb:split:{ticker}",
                notes=f"Split {format(ratio.normalize(), 'f')}:1 (XTB transfer)",
                ratio=ratio,
            )
            split_rows.append(
                self._asset(row, {"Ticker": ticker, "Category": old[0].category})
            )
        return split_rows, split_issues

    # --- Expectations ---

    def _open_positions(self, workbook: Any) -> dict[str, Decimal]:
        held: dict[str, Decimal] = {}
        for _, cells in _table(workbook, _OPEN_SHEET, "Product", also="Volume"):
            side = _text(cells.get("Type")).upper()
            volume = _decimal(cells.get("Volume"))
            ticker = _text(cells.get("Ticker"))
            if side not in {"BUY", "SELL"} or volume is None or not ticker:
                continue
            signed = volume if side == "BUY" else -volume
            held[ticker] = held.get(ticker, Decimal("0")) + signed
        return held

    def _closed_profit(self, workbook: Any) -> Decimal | None:
        for _, cells in _table(workbook, _CLOSED_SHEET, "Instrument"):
            if _text(cells.get("Instrument")).lower() == "profit/loss":
                return _money(cells.get("Profit/Loss"))
        return None


def _stable_refs(rows: list[ParsedRow]) -> list[ParsedRow]:
    """Every ref gets the row's own timestamp: the source's ids may be reused or
    anonymised, but id + time identifies a row the same way in any export, so a
    later export of the same account still deduplicates. Rows identical even
    so (same id and time) get an occurrence index."""
    seen: Counter[str] = Counter()
    result: list[ParsedRow] = []
    for row in rows:
        ref = f"{row.external_ref}:{row.operation_date.isoformat()}"
        seen[ref] += 1
        if seen[ref] > 1:
            ref = f"{ref}:{seen[ref]}"
        result.append(replace(row, external_ref=ref))
    return result
