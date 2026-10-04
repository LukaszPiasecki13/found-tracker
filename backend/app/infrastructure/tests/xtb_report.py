"""Builds an XTB account report (.xlsx) in memory, laid out like the real one,
so tests need neither a bank file nor personal data."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook

CASH_HEADER = (
    "Type",
    "Instrument",
    "Ticker",
    "Category",
    "Time",
    "Amount",
    "ID",
    "Comment",
    "Product",
    "Position ID",
)


def _at(day: int, hour: int = 10) -> datetime:
    """A naive datetime, as `openpyxl` reads date cells (the report is UTC)."""
    return datetime(2026, 1, day, hour)


def cash_row(
    type_: str,
    when: datetime,
    amount: float,
    *,
    id_: int | str = "",
    ticker: str = "",
    comment: str = "",
    position_id: str = "",
) -> tuple[Any, ...]:
    instrument = ticker.split(".")[0] if ticker else ""
    return (
        type_,
        instrument,
        ticker,
        "STOCK" if ticker else "",
        when,
        amount,
        id_,
        comment,
        "My Trades",
        position_id,
    )


# A complete, ledger-valid history: every sell has its buy, cash never < 0.
COMPLETE_ROWS: list[tuple[Any, ...]] = [
    cash_row("Deposit", _at(10), 5000, id_=1, comment="Deposit"),
    cash_row(
        "Stock purchase",
        _at(11),
        -400,
        id_=2,
        ticker="DNP.PL",
        comment="OPEN BUY 10 @ 40.000",
        position_id="100",
    ),
    cash_row(
        "Stock purchase",
        _at(12),
        -200,
        id_=3,
        ticker="ETFBM40TR.PL",
        comment="OPEN BUY 2 @ 100.000",
        position_id="101",
    ),
    cash_row("Free funds interest", _at(15), 0.5, id_=4, comment="Interest 2026-01"),
    cash_row("Free funds interest tax", _at(15, 11), -0.1, id_=5, comment="Tax"),
    cash_row(
        "Stock sell",
        _at(20),
        200,
        id_=6,
        ticker="DNP.PL",
        comment="CLOSE BUY 4/10 @ 50.000",
        position_id="100",
    ),
    cash_row(
        "Dividend",
        _at(22),
        6,
        id_=8,
        ticker="DNP.PL",
        comment="DNP.PL PLN 1.0000/ SHR",
        position_id="100",
    ),
    cash_row("Withdrawal", _at(25), -300, id_=7, comment="Withdrawal"),
]
COMPLETE_CASH_TOTAL = Decimal("4306.40")
COMPLETE_OPEN_POSITIONS: dict[str, float] = {"DNP.PL": 6, "ETFBM40TR.PL": 2}
COMPLETE_TOTAL_DEPOSITED = Decimal("4700")


def closed_row(
    ticker: str,
    volume: float,
    opened: datetime,
    closed: datetime,
    position_id: str,
    comment: str = "",
) -> tuple[Any, ...]:
    """A "Closed Positions" row (the columns the parser reads)."""
    return (
        ticker.split(".")[0],
        ticker,
        "STOCK",
        "BUY",
        volume,
        None,
        opened,
        closed,
        "xStation5",
        position_id,
        comment,
    )


def transfer_out_row(
    ticker: str, volume: float, opened: datetime, closed: datetime, position_id: str
) -> tuple[Any, ...]:
    """What XTB writes when a split closes an old position."""
    row = list(closed_row(ticker, volume, opened, closed, position_id))
    row[8] = "Correction"
    row[10] = "STC Transfer Out"
    return tuple(row)


def build_xtb_report(
    cash_rows: list[tuple[Any, ...]] | None = None,
    *,
    cash_total: float | None = float(COMPLETE_CASH_TOTAL),
    open_positions: dict[str, float] | None = None,
    closed_profit: float | None = 12.5,
    include_cash_sheet: bool = True,
    closed_rows: list[tuple[Any, ...]] | None = None,
    open_lots: list[tuple[str, str, float, datetime]] | None = None,
) -> bytes:
    """An .xlsx like XTB's export. `cash_total` is the "Total" row (`None`
    leaves it out); `open_positions` maps tickers to held volume."""
    rows = COMPLETE_ROWS if cash_rows is None else cash_rows
    held = COMPLETE_OPEN_POSITIONS if open_positions is None else open_positions
    workbook = Workbook()
    closed = workbook.active
    assert closed is not None
    closed.title = "Closed Positions"
    closed.append(("Account number", 1111111))
    closed.append(("Closed Positions", ""))
    closed.append(("Date from (UTC)", _at(1)))
    closed.append(("Date to (UTC)", _at(28)))
    closed.append(
        (
            "Instrument",
            "Ticker",
            "Category",
            "Type",
            "Volume",
            "Profit/Loss",
            "Open Time (UTC)",
            "Close Time (UTC)",
            "Close Origin",
            "Position ID",
            "Comment",
        )
    )
    for closed_position in closed_rows or []:
        closed.append(closed_position)
    if closed_profit is not None:
        closed.append(("Profit/loss", None, None, None, None, closed_profit))

    if include_cash_sheet:
        cash = workbook.create_sheet("Cash Operations")
        cash.append(("Account number", 111111))
        cash.append(("Cash Operations", ""))
        cash.append(("Date from (UTC)", _at(1)))
        cash.append(("Date to (UTC)", _at(28)))
        cash.append(CASH_HEADER)
        for row in rows:
            cash.append(row)
        if cash_total is not None:
            cash.append(("Total", None, None, None, None, cash_total))

    opened = workbook.create_sheet("Open Positions")
    opened.append(("Account number", 111111))
    opened.append(("Open Positions", ""))
    opened.append(("Data as of report generated", _at(28)))
    opened.append(("Product", "Metric", "Amount", "Currency"))
    opened.append(("My Trades", "Open position value", 100, "PLN"))
    opened.append((None,))
    opened.append(("Note", "Summary values and open positions are shown as of now"))
    opened.append(
        (
            "Product",
            "Instrument/Position",
            "Ticker",
            "Category",
            "Type",
            "Volume",
            "Open time (UTC)",
        )
    )
    for position_id, ticker, volume, open_time in open_lots or []:
        opened.append(("My Trades", position_id, ticker, "", "BUY", volume, open_time))
    for ticker, volume in held.items():
        opened.append(("My Trades", ticker.split(".")[0], ticker, "STOCK", "", volume))
        opened.append(("My Trades", 999, ticker, "", "BUY", volume))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
